"""Generate the demo documents served by the mock portal.

Deterministic on purpose: the demo numbers in roadmap.md assume these exact
files. Re-run to restore them after a demo reset.

Vendor codes and names are real, taken from Oracle EBS suppliers that have 2026
purchase orders at operating unit 152, and the part numbers are items actually
bought from them. Real master data means DBQUERY_MODE=mock and DBQUERY_MODE=real
answer the same questions the same way.

Two portal sections, so the agent has a routing decision to make:

  SCM Documents          supplier_001.xlsx  invoice, supplier code present
                         quality_002.xlsx   vendor NAME only -> lookup by name
                         debit_003.csv      debit note, supplier code present
                         unknown_004.txt    unstructured prose -> manual review

  ESG Questionnaires     esg_3707.xlsx      complete, supplier code present
                         esg_9414.xlsx      complete, supplier code present
                         esg_draft_005.xlsx unanswered sections -> manual review

Run inside the mcp-worker image (it already has openpyxl):
  docker run --rm -v <repo>/mock-portal/MockData:/out \
      -v <repo>/scripts:/scripts:ro \
      agentic-document-demo-mcp-worker python /scripts/gen_mock_data.py /out
"""

import csv
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

OUT = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "/out")

LABEL = Font(bold=True)
TITLE = Font(bold=True, size=14)


def _autosize(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _fields(ws, pairs, start_row=3):
    for i, (k, v) in enumerate(pairs, start=start_row):
        ws.cell(row=i, column=1, value=k).font = LABEL
        ws.cell(row=i, column=2, value=v)
    return start_row + len(pairs)


def _title(ws, text, span="A1:D1"):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws.merge_cells(span)
    ws["A1"].alignment = Alignment(horizontal="left")


# ---------------------------------------------------------------------------
# SCM documents
# ---------------------------------------------------------------------------

def supplier_invoice(path):
    """Well-formed supplier invoice. Everything mapping needs is present."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    _title(ws, "SUPPLIER INVOICE")

    _fields(ws, [
        ("InvoiceNo", "INV-2026-0912"),
        ("InvoiceDate", "2026-09-12"),
        ("SupplierCode", "3707"),
        ("SupplierName", "華碩電腦股份有限公司"),
        ("Currency", "TWD"),
        ("Department", "SCM"),
    ])

    ws.cell(row=10, column=1, value="Line Items").font = LABEL
    for c, name in enumerate(["PartNo", "Description", "Qty", "UnitPrice", "Amount"], start=1):
        ws.cell(row=11, column=c, value=name).font = LABEL

    rows = [
        ("05C531-103530", "THERMISTOR NTC.10K 3% SMD 0603", 5000, 1.8, 9000),
        ("04-888-105003", "C/C.1uF.10V.10% X5R SMD 0402", 20000, 0.35, 7000),
    ]
    for r, row in enumerate(rows, start=12):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)

    ws.cell(row=15, column=4, value="Total").font = LABEL
    ws.cell(row=15, column=5, value=16000)

    _autosize(ws, {"A": 18, "B": 34, "C": 10, "D": 12, "E": 12})
    wb.save(path)


def quality_report(path):
    """Quality report naming a vendor but carrying no supplier code.

    勝宏科技 is a real two-entity group in EBS (Huizhou and Thailand). Whether
    this resolves depends on scripts/scenario.py, which is the Phase 7 contrast.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "IQC Report"
    _title(ws, "INCOMING QUALITY INSPECTION REPORT", "A1:C1")

    _fields(ws, [
        ("Vendor", "勝宏科技"),
        ("Lot", "LOT-26-99123"),
        ("InspectionDate", "2026-09-12"),
        ("Inspector", "QA-07"),
        ("SampleSize", 50),
        ("Defects", 6),
        ("Inspection Result", "FAIL"),
        ("Department", "SCM"),
    ])

    ws.cell(row=12, column=1, value="Remark").font = LABEL
    ws.cell(row=12, column=2, value="Solder pad oxidation found on 6 of 50 sampled units.")

    _autosize(ws, {"A": 20, "B": 56, "C": 12})
    wb.save(path)


def debit_note(path):
    """Flat CSV debit note. Supplier code present, so mapping is direct."""
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["supplier", "supplier_name", "reason", "amount", "currency", "debit_no", "department"])
        w.writerow(["3385", "聯強國際股份有限公司", "late delivery", 25000, "TWD", "DN-2026-0093", "SCM"])


def unknown_notice(path):
    """Unstructured prose with no dependable fields. Expected to reach manual review."""
    path.write_text(
        "至上電子股份有限公司 業務部 收\n"
        "\n"
        "We have identified abnormal damage during incoming inspection of the\n"
        "shipment received this week.\n"
        "\n"
        "The affected cartons arrived with visible crushing on the outer\n"
        "packaging. Pending your reply we are holding the shipment in the\n"
        "quarantine area and have not booked it into stock.\n"
        "\n"
        "Reference: LOT-26-88881\n"
        "\n"
        "請盡快回覆處理方式。\n",
        encoding="utf-8",
    )


# ---------------------------------------------------------------------------
# ESG questionnaires
# ---------------------------------------------------------------------------

SECTIONS = [
    ("E1", "Greenhouse gas inventory (Scope 1 & 2)"),
    ("E2", "Renewable electricity ratio"),
    ("E3", "Hazardous substance management (RoHS / REACH)"),
    ("S1", "Labour rights and working hours"),
    ("S2", "Occupational health and safety"),
    ("G1", "Anti-corruption policy"),
    ("G2", "Conflict minerals declaration"),
]


def esg_survey(path, code, name, answers, submitted="2026-09-12", status="Submitted"):
    wb = Workbook()
    ws = wb.active
    ws.title = "ESG Survey"
    _title(ws, "SUPPLIER ESG QUESTIONNAIRE 2026", "A1:C1")

    next_row = _fields(ws, [
        ("SurveyNo", f"ESG-2026-{code}"),
        ("SupplierCode", code),
        ("SupplierName", name),
        ("SubmittedOn", submitted),
        ("Status", status),
        ("Department", "ESG"),
    ])

    header = next_row + 1
    for c, label in enumerate(["Section", "Question", "Response"], start=1):
        ws.cell(row=header, column=c, value=label).font = LABEL

    for i, (section, question) in enumerate(SECTIONS, start=header + 1):
        ws.cell(row=i, column=1, value=section)
        ws.cell(row=i, column=2, value=question)
        ws.cell(row=i, column=3, value=answers.get(section, ""))

    _autosize(ws, {"A": 12, "B": 52, "C": 22})
    wb.save(path)


def esg_complete_a(path):
    esg_survey(path, "3707", "華碩電腦股份有限公司", {
        "E1": "Verified, ISO 14064-1",
        "E2": "42%",
        "E3": "Compliant",
        "S1": "Compliant",
        "S2": "ISO 45001 certified",
        "G1": "Policy published",
        "G2": "CMRT submitted",
    })


def esg_complete_b(path):
    esg_survey(path, "9414", "瀚宇博德科技(江陰)有限公司", {
        "E1": "Verified, ISO 14064-1",
        "E2": "18%",
        "E3": "Compliant",
        "S1": "Compliant",
        "S2": "ISO 45001 certified",
        "G1": "Policy published",
        "G2": "CMRT submitted",
    })


def esg_draft(path):
    """Half-finished return with no supplier code. Expected to reach manual review."""
    esg_survey(path, "", "麗臺科技", {
        "E1": "In progress",
        "E3": "Compliant",
        "S1": "Compliant",
    }, submitted="", status="Draft")


# ---------------------------------------------------------------------------

BUILDERS = [
    ("scm", "supplier_001.xlsx", supplier_invoice),
    ("scm", "quality_002.xlsx", quality_report),
    ("scm", "debit_003.csv", debit_note),
    ("scm", "unknown_004.txt", unknown_notice),
    ("esg", "esg_3707.xlsx", esg_complete_a),
    ("esg", "esg_9414.xlsx", esg_complete_b),
    ("esg", "esg_draft_005.xlsx", esg_draft),
]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for section, name, build in BUILDERS:
        target = OUT / name
        build(target)
        print(f"[{section}] wrote {target.name} ({target.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
