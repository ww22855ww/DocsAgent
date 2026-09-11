"""Turn a staged file into plain text the LLM can classify.

Output is deliberately flat and label-preserving: the classifier needs to see
"SupplierCode: V00123", not a reconstructed table. Content is capped so one
oversized document cannot blow the model's context budget.
"""

from __future__ import annotations

import csv
import io
import os

from openpyxl import load_workbook

MAX_CHARS = 6000
MAX_XLSX_ROWS = 200
MAX_CSV_ROWS = 200


def _truncate(text: str) -> tuple[str, bool]:
    if len(text) <= MAX_CHARS:
        return text, False
    return text[:MAX_CHARS] + "\n... [truncated]", True


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_xlsx(path: str) -> tuple[str, dict]:
    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = []
    lines: list[str] = []

    for ws in wb.worksheets:
        sheets.append(ws.title)
        lines.append(f"# Sheet: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_XLSX_ROWS:
                lines.append("... [more rows omitted]")
                break
            cells = [_cell(c) for c in row]
            while cells and not cells[-1]:
                cells.pop()
            if not cells:
                continue
            # Two-column rows are label/value pairs; keep that shape readable.
            if len(cells) == 2:
                lines.append(f"{cells[0]}: {cells[1]}")
            else:
                lines.append(" | ".join(cells))
    wb.close()

    return "\n".join(lines), {"sheets": sheets, "sheet_count": len(sheets)}


def parse_csv(path: str) -> tuple[str, dict]:
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows = []
        for i, row in enumerate(reader):
            if i >= MAX_CSV_ROWS:
                rows.append(["... [more rows omitted]"])
                break
            rows.append(row)

    if not rows:
        return "", {"columns": [], "row_count": 0}

    header = rows[0]
    lines = [" | ".join(header)]
    for row in rows[1:]:
        lines.append(" | ".join(row))

    # Also emit the first data row as labelled pairs; short CSVs classify far
    # more reliably when the header is attached to the value.
    if len(rows) > 1 and len(rows[1]) == len(header):
        lines.append("")
        lines.append("# First record")
        for k, v in zip(header, rows[1]):
            lines.append(f"{k}: {v}")

    return "\n".join(lines), {"columns": header, "row_count": max(0, len(rows) - 1)}


def parse_text(path: str) -> tuple[str, dict]:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()
    return content, {"line_count": content.count("\n") + 1}


def extract(path: str) -> dict:
    """Parse one file. Unsupported types return an error entry rather than raising,
    so a single bad file cannot fail a whole batch."""
    name = os.path.basename(path)
    ext = os.path.splitext(name)[1].lower()

    try:
        if ext == ".xlsx":
            content, meta = parse_xlsx(path)
        elif ext == ".csv":
            content, meta = parse_csv(path)
        elif ext in (".txt", ".md", ".log"):
            content, meta = parse_text(path)
        else:
            return {
                "filename": name,
                "content": "",
                "metadata": {},
                "error": f"Unsupported file type: {ext or '(none)'}",
            }
    except Exception as exc:  # noqa: BLE001 - surfaced to the agent as data
        return {"filename": name, "content": "", "metadata": {}, "error": f"{type(exc).__name__}: {exc}"}

    content, truncated = _truncate(content)
    meta = {**meta, "extension": ext, "bytes": os.path.getsize(path), "truncated": truncated}
    return {"filename": name, "content": content, "metadata": meta, "error": None}
