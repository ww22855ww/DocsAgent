"""Excel export of a finished task.

Reads back what was actually committed (archives and manual_reviews), not what
the agent believes it did, so the workbook and the database can never disagree.
"""

from __future__ import annotations

import logging
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from .common import ToolError, db

log = logging.getLogger("mcp-worker.excel")

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="26334D")
TITLE_FONT = Font(bold=True, size=13)


def _write_sheet(ws, headers: list[str], rows: list[list], widths: list[int]) -> None:
    for c, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"


def write_excel(task_id: str, filename: str | None = None) -> dict:
    """Write the result workbook for one task into the output directory."""
    if not task_id or not task_id.strip():
        raise ToolError("task_id is required.")
    task_id = task_id.strip()

    with db() as conn, conn.cursor() as cur:
        cur.execute(
            """
            SELECT filename, category, confidence, document_no, supplier_code,
                   supplier_name, part_no, archived_at, archived_path
            FROM archives WHERE task_id = %s ORDER BY id
            """,
            (task_id,),
        )
        archived = cur.fetchall()

        cur.execute(
            """
            SELECT filename, reason, summary, notified, notify_detail, created_at
            FROM manual_reviews WHERE task_id = %s ORDER BY id
            """,
            (task_id,),
        )
        reviews = cur.fetchall()

    if not archived and not reviews:
        raise ToolError(f"No archived documents or manual reviews recorded for {task_id}.")

    wb = Workbook()

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Agentic Document Processing — Task Result"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")

    by_category: dict[str, int] = {}
    for row in archived:
        key = row["category"] or "(unclassified)"
        by_category[key] = by_category.get(key, 0) + 1

    summary_rows = [
        ("Task", task_id),
        ("Documents processed", len(archived) + len(reviews)),
        ("Archived automatically", len(archived)),
        ("Manual review", len(reviews)),
        ("Notifications sent", sum(1 for r in reviews if r["notified"])),
        ("", ""),
        ("By category", ""),
        *[(f"  {k}", v) for k, v in sorted(by_category.items())],
    ]
    for i, (k, v) in enumerate(summary_rows, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=not str(k).startswith("  "))
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 40

    # Archived
    _write_sheet(
        wb.create_sheet("Archived"),
        ["File", "Category", "Confidence", "Document No", "Supplier Code",
         "Supplier Name", "Part No", "Archived At"],
        [[r["filename"], r["category"],
          float(r["confidence"]) if r["confidence"] is not None else None,
          r["document_no"], r["supplier_code"], r["supplier_name"], r["part_no"],
          r["archived_at"].strftime("%Y-%m-%d %H:%M:%S") if r["archived_at"] else None]
         for r in archived],
        [22, 18, 11, 16, 14, 24, 14, 20],
    )

    # Manual review
    _write_sheet(
        wb.create_sheet("Manual Review"),
        ["File", "Reason", "Summary", "Notified", "Detail", "Raised At"],
        [[r["filename"], r["reason"], r["summary"], "Yes" if r["notified"] else "No",
          r["notify_detail"],
          r["created_at"].strftime("%Y-%m-%d %H:%M:%S") if r["created_at"] else None]
         for r in reviews],
        [22, 42, 46, 10, 46, 20],
    )

    os.makedirs(config.OUTPUT_DIR, exist_ok=True)
    name = (filename or "result.xlsx").strip()
    if "/" in name or "\\" in name or not name.endswith(".xlsx"):
        raise ToolError("filename must be a plain .xlsx name.")

    path = os.path.join(config.OUTPUT_DIR, name)
    wb.save(path)
    log.info("wrote %s (%d archived, %d manual review)", path, len(archived), len(reviews))

    return {
        "status": "written",
        "path": path,
        "filename": name,
        "archived": len(archived),
        "manual_review": len(reviews),
        "bytes": os.path.getsize(path),
    }
